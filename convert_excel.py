import openpyxl
import json

# Load the Excel file
wb = openpyxl.load_workbook('one manolo manual fare input_03092025.xlsx')

# Print sheet names to understand the structure
print("Sheet names:", wb.sheetnames)

# Read each sheet
data = {}
for sheet_name in wb.sheetnames:
    sheet = wb[sheet_name]
    sheet_data = []
    
    # Get headers from first row
    headers = []
    for cell in sheet[1]:
        headers.append(cell.value)
    
    print(f"\nSheet: {sheet_name}")
    print(f"Headers: {headers}")
    
    # Read data rows
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if any(cell is not None for cell in row):  # Skip empty rows
            row_dict = {}
            for i, value in enumerate(row):
                if i < len(headers):
                    row_dict[headers[i]] = value
            sheet_data.append(row_dict)
    
    data[sheet_name] = sheet_data
    print(f"Rows: {len(sheet_data)}")

# Save to JSON file
with open('fare_metrics.json', 'w', encoding='utf-8') as f:
    json.dump(data, f, indent=2, default=str)

print("\nData saved to fare_metrics.json")
